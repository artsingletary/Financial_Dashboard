import os
from openpyxl import load_workbook
from config import (
    WORKBOOK_FOLDER,
    WORKBOOK_NAME,
    DASHBOARD_SHEET,
    MARKET_SHEET,
    HISTORY_SHEET,
    ACTIVITY_SHEET,
    HOLDINGS_SHEET,
    HOLDINGS_COL
)   

def open_workbook():
    file = os.path.join(WORKBOOK_FOLDER, WORKBOOK_NAME)
    if not os.path.exists(file):
        raise FileNotFoundError(f"Workbook not found at {file}. Please ensure the file exists.")
    wb = load_workbook(file)
    return wb, file

def get_market_sheet(wb):
    return wb[MARKET_SHEET]

def get_dashboard_sheet(wb):
    return wb[DASHBOARD_SHEET]

def get_history_sheet(wb):
    for name in wb.sheetnames:
        if name.strip().lower() == "history":
            return wb[name]
    return wb.create_sheet(HISTORY_SHEET)

def get_activity_sheet(wb):
    return wb[ACTIVITY_SHEET]

def get_holdings_sheet(wb):
    return wb[HOLDINGS_SHEET]

def save_workbook(wb, file):
    wb.save(file)