import os
from openpyxl.styles import PatternFill


# File Settings
WORKBOOK_FOLDER = os.path.expanduser("~/OneDrive/Active")
WORKBOOK_NAME = "Financial Dashboard.xlsx"

# Worksheet Names
DASHBOARD_SHEET = "Dashboard"
MARKET_SHEET = "Market Data"
HISTORY_SHEET = "History"
ACTIVITY_SHEET = "Activity Log"
HOLDINGS_SHEET = "Holdings"
TSP_PRICE_SHEET = "TSP Prices"

HOLDINGS_COL = {
    "account_type": "A",
    "symbol": "B",
    "expense_ratio": "C",
    "description": "D",
    "shares": "E",
    "price": "F",
    "market_value": "G",
    "source": "H",
    "asset_class": "I",
    "allocation_category": "J",
    "account": "K",
    "owner": "L",
    "tax_status": "M",
    "target": "N",
    "notes": "O", 
    "raw_price": "P"
    }

# Dashboard Cells Mapping
DASHBOARD_CELLS = {
   "Dow": {
      "price": "B25",
      "status": "C25"
   },
   "S&P 500": {
      "price": "B26",
      "status": "C26"
   },
   "Nasdaq": {
      "price": "B27",
      "status": "C27"
   }
}


# Dashboard Thresholds
GREEN_THRESHOLD = -0.01  # Within 1% of record high
YELLOW_THRESHOLD = -0.03  # Within 3% of record high

# Dashboard Colors
GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

# Markets 
MARKETS = {
   "Dow": {
      "ticker": "^DJI",
      "row": 2
   },
   "S&P 500": {
     "ticker": "^GSPC",
       "row": 3
   },
   "Nasdaq": { 
   "ticker": "^IXIC",
   "row": 4
   }
}